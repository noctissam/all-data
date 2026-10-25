# -*- coding: utf-8 -*-
"""
Build the name -> CAS lookup the workbook is written from.

The workbook is the source of the name list, not a separate file of names: the names in
the sheets are the names that need numbers, and a name that the extraction produced but
never reached a block needs none. Reading them out of the built workbook also means the
lookup follows whatever block definition the readers use, so the three cannot drift apart.

Only names whose status is a success contribute a CAS. An untranslatable name, a name
PubChem does not index, and a name whose only match was a neighbouring compound are three
different answers and none of them is a CAS; each leaves the cell blank with a reason.

Run: python collect_names.py [--xlsx PATH] [--out PATH]
"""

import argparse
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import blocks  # noqa: E402

REOUT = os.path.abspath(os.path.join(HERE, "..", "out", "re整理"))
CACHE = os.path.join(REOUT, "pubchem_cache.json")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=os.path.join(REOUT, "re整理.xlsx"))
    ap.add_argument("--out", default=os.path.join(REOUT, "cas_by_name.json"))
    a = ap.parse_args()

    cache = json.load(open(CACHE, encoding="utf-8")) if os.path.exists(CACHE) else {}
    ok = {k: v[0] for k, v in cache.items() if v[0] and str(v[2]).startswith("ok")}
    print("cached CAS usable: %d" % len(ok))

    import openpyxl
    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    want = {}
    for sn in wb.sheetnames:
        if sn == "00_索引":
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        for start, stop in blocks.find_data_blocks(rows):
            for row in rows[start + 1:stop]:
                a0 = row[0] if row else None
                if a0 is None:
                    continue
                name = str(a0).strip()
                if not name:
                    continue
                want[name] = want.get(name, 0) + 1
    wb.close()
    print("names in workbook: %d" % len(want))

    hit = {n: ok[n] for n in want if n in ok}
    print("names with a resolved CAS: %d  (%.0f%% of names, %d of %d rows)"
          % (len(hit), 100.0 * len(hit) / max(len(want), 1),
             sum(want[n] for n in hit), sum(want.values())))

    json.dump({"cas": hit, "rows": want},
              open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("-> %s" % a.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())