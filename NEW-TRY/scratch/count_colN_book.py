# -*- coding: utf-8 -*-
"""Count the 列N column headers that survive into the delivered workbook."""
import collections
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月\re整理.xlsx"
wb = openpyxl.load_workbook(P, data_only=True, read_only=True)

hits = collections.Counter()
samples = collections.defaultdict(list)
for sn in wb.sheetnames:
    if sn == "00_索引":
        continue
    ws = wb[sn]
    head = None
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        first = str(cells[0] or "").strip()
        if "基础数据" in first:
            head = None
            continue
        if head is None:
            if first.startswith("化合物名称"):
                head = [str(x or "") for x in cells]
                for j, h in enumerate(head):
                    if re.search(r"浓度\s*列\d+$", h) or re.search(r"阈值\s*列\d+$", h):
                        hits[(sn[:22], h)] += 1
                        samples[sn[:22]].append(h)
            continue

print("交付工作簿中仍为「列N」的列标题：%d 列 / %d 表" % (sum(hits.values()), len(samples)))
for sn, cols in samples.items():
    print("  %-26s %s" % (sn, ", ".join(sorted(set(cols)))))
wb.close()