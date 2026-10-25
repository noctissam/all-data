# -*- coding: utf-8 -*-
"""Check 清洗后数据.xlsx after the threshold/OAV columns were folded in."""
import collections
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月\清洗后数据.xlsx"
wb = openpyxl.load_workbook(P, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
print("sheet:", wb.sheetnames[0])
print("表头:", [c.value for c in next(ws.iter_rows(max_row=1))])
rows = list(ws.iter_rows(min_row=2, values_only=True))
print("数据行:", len(rows))

num = lambda v: isinstance(v, (int, float))  # noqa: E731
bad_pm = sum(1 for r in rows if r[5] is not None and "±" in str(r[5]))
bad_al = sum(1 for r in rows if isinstance(r[5], str) and re.search(r"\d[A-Za-z]$", r[5]))
bad_cm = sum(1 for r in rows if isinstance(r[5], str) and "," in r[5])
print("仍含 ± / 尾字母 / 千分位逗号 : %d / %d / %d" % (bad_pm, bad_al, bad_cm))

kind = collections.Counter(r[3] for r in rows)
print("数据类型:", dict(kind))
nkind = collections.Counter()
for r in rows:
    nkind[(r[3], "数值" if num(r[5]) else "文本")] += 1
for k, v in sorted(nkind.items()):
    print("   %-10s %-6s %6d" % (k[0], k[1], v))

texts = collections.Counter(str(r[5]) for r in rows if not num(r[5]))
print("非数值取值:", texts.most_common(10))

print("空值: 文献名 %d  化合物 %d  数据类型 %d  数值 %d"
      % (sum(1 for r in rows if not r[0]), sum(1 for r in rows if not r[1]),
         sum(1 for r in rows if not r[3]), sum(1 for r in rows if r[5] is None)))
print("不同文献 %d   不同样品 %d" % (len({r[0] for r in rows}), len({r[4] for r in rows})))

print("\n前 10 行:")
for r in rows[:10]:
    print("   %-30s %-22s %-11s %-5s %-14s %s"
          % (str(r[0])[:30], str(r[1])[:22], str(r[2])[:11], str(r[3]), str(r[4])[:14], r[5]))
wb.close()