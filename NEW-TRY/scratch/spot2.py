# -*- coding: utf-8 -*-
"""Spot-check the clean file backwards: for each sample, find the raw cell it came from."""
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月"

src = openpyxl.load_workbook(os.path.join(ROOT, "re整理.xlsx"), data_only=True)
cln = openpyxl.load_workbook(os.path.join(ROOT, "清洗后数据.xlsx"), data_only=True,
                             read_only=True)
rows = list(cln[cln.sheetnames[0]].iter_rows(min_row=2, values_only=True))

# build (name, sample-ish, raw) for every concentration cell in the source
raw = {}
for sn in src.sheetnames:
    if sn == "00_索引":
        continue
    ws = src[sn]
    head = None
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        first = str(cells[0] or "").strip()
        if "基础数据" in first:
            head = None
            continue
        if head is None:
            if first.startswith("化合物名称"):
                head = [str(x or "") for x in cells]
            continue
        conc = [j for j, h in enumerate(head) if h.startswith("浓度")]
        for j in conc:
            if j < len(cells) and cells[j] not in (None, ""):
                raw.setdefault(str(first), set()).add(str(cells[j]).strip())

picks = [rows[0], rows[2000], rows[6000], rows[9000], rows[-1]]
print("%-26s %-22s %-16s %s" % ("化合物", "样品", "清洗后", "源原始"))
print("-" * 96)
ok = miss = 0
for r in picks:
    name, sample, val = r[1], r[3], r[4]
    cands = raw.get(name, set())
    hit = None
    for c in cands:
        head = re.split(r"[±]", c, 1)[0].strip()
        head = re.sub(r"[A-Za-z]+$", "", head)
        head = head.replace(",", "").replace(" ", "").replace("\u00a0", "")
        try:
            if abs(float(head) - float(val)) < 1e-9:
                hit = c
                break
        except (TypeError, ValueError):
            if head == str(val):
                hit = c
                break
    if hit:
        ok += 1
    else:
        miss += 1
    print("%-26s %-22s %-16s %s"
          % (str(name)[:26], str(sample)[:22], val, hit or ("候选: " + str(list(cands)[:2]))))
print("-" * 96)
print("抽查 5 行: 命中源值 %d, 未命中 %d" % (ok, miss))
src.close()
cln.close()