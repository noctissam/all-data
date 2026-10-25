# -*- coding: utf-8 -*-
"""Show the 基础数据 block layout of the first few sheets."""
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月"

wb = openpyxl.load_workbook(os.path.join(ROOT, "re整理.xlsx"), data_only=True)
print("sheets:", len(wb.sheetnames))
print("names:", wb.sheetnames[:6], "...")
print()

for sn in wb.sheetnames[1:4]:
    ws = wb[sn]
    print("=" * 72)
    print("sheet:", sn, " dims:", ws.max_row, "x", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i > 22:
            break
        cells = ["" if x is None else str(x) for x in row[:7]]
        if not any(cells):
            continue
        print("  %3d | %s" % (i, " | ".join(c[:26] for c in cells)))
wb.close()