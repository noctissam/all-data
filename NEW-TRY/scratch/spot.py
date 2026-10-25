# -*- coding: utf-8 -*-
"""Spot-check: does 清洗后数据 agree with what the source workbook actually printed?"""
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月"

src = openpyxl.load_workbook(os.path.join(ROOT, "re整理.xlsx"), data_only=True)
cln = openpyxl.load_workbook(os.path.join(ROOT, "清洗后数据.xlsx"), data_only=True,
                             read_only=True)
cln_rows = list(cln[cln.sheetnames[0]].iter_rows(min_row=2, values_only=True))

CHECKS = [
    ("01 2  Characterization of key o", "ethyl hexanoate"),
    ("12 11不同等级浓香型白酒基酒挥发性风味物质差异分��",
     "乙酸乙酯"),
    ("20 19  浓香、酱香、清香型白酒挥发性风味的", "乙酸乙酯"),
]

for sn, want in CHECKS:
    ws = src[sn]
    head = None
    print("=" * 74)
    print("源 sheet:", sn)
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        first = str(cells[0] or "").strip()
        if "基础数据" in first:
            head = None
            continue
        if head is None:
            if first.startswith("化合物名称"):
                head = [str(x or "") for x in cells]
                print("  表头:", " | ".join(h[:22] for h in head if h.strip()))
            continue
        if first == want:
            print("  源原始值 :", " | ".join(
                "" if cells[j] is None else str(cells[j])[:26]
                for j in range(min(len(head), 9))))
            for r in cln_rows:
                if r[1] == want and r[0].startswith(
                        str(ws.cell(row=1, column=1).value or "")[:12]):
                    print("  清洗后   : 样品=%-24s 浓度=%s" % (r[3], r[4]))
            break
src.close()
cln.close()