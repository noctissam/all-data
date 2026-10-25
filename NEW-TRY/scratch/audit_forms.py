# -*- coding: utf-8 -*-
"""Audit every raw cell the 浓度 columns hold, to see which value forms exist.

The point is to find notation the cleaner does not recognise - a "mean ± spread" written
some other way would otherwise pass through as text and sit in a numeric column.
"""
import collections
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
SRC = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月\NEW-TRY\out\re整理\re整理.xlsx"

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))
from make_clean import clean_value  # noqa: E402

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
forms = collections.Counter()
examples = {}
for sn in wb.sheetnames:
    if sn == "00_索引":
        continue
    ws = wb[sn]
    head = None
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        first = str(cells[0] or "").strip()
        if "基础数据" in first:
            head = None
            continue
        if head is None:
            if first.startswith("化合物名称"):
                head = [str(x or "") for x in cells]
            continue
        conc = [j for j, h in enumerate(head) if h.startswith("浓度")]
        if not conc:
            continue
        for j in conc:
            if j >= len(cells):
                continue
            v = cells[j]
            if v is None or str(v).strip() == "":
                continue
            s = str(v).strip()
            out = clean_value(v)
            if isinstance(out, float):
                # did the raw text actually contain a spread notation we removed?
                if re.search(r"[±~〜∓]", s) or re.search(r"[A-Za-z]", s) or "," in s:
                    key = "含波动/字母/逗号 -> 已取定量值"
                else:
                    key = "本来就是单个数"
            else:
                key = "清洗后仍是文本"
            forms[key] += 1
            if key == "清洗后仍是文本":
                examples.setdefault(s, 0)
                examples[s] += 1
            elif key == "含波动/字母/逗号 -> 已取定量值":
                pass
wb.close()

for k, v in forms.most_common():
    print("%-34s %6d" % (k, v))
print()
print("清洗后仍为文本的取值（去重计数，共 %d 种）:" % len(examples))
for s, n in sorted(examples.items(), key=lambda x: -x[1])[:40]:
    print("   %-24s x%d" % (repr(s), n))