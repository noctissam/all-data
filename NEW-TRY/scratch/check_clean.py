# -*- coding: utf-8 -*-
"""Check 清洗后数据.xlsx: did the cleaning actually clean?"""
import collections
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月\清洗后数据.xlsx"
wb = openpyxl.load_workbook(P, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
print("sheet:", wb.sheetnames, " header:",
      [c.value for c in next(ws.iter_rows(max_row=1))])

rows = list(ws.iter_rows(min_row=2, values_only=True))
print("data rows:", len(rows))

bad_pm = [r for r in rows if r[4] is not None and "±" in str(r[4])]
bad_alpha = [r for r in rows if isinstance(r[4], str)
             and re.search(r"\d[A-Za-z]$", str(r[4]))]
bad_comma = [r for r in rows if isinstance(r[4], str) and "," in str(r[4])]
print("仍含 ±        :", len(bad_pm))
print("数字带尾字母   :", len(bad_alpha))
print("仍含千分位逗号 :", len(bad_comma))

kind = collections.Counter()
texts = collections.Counter()
for r in rows:
    v = r[4]
    if v is None:
        kind["空"] += 1
    elif isinstance(v, (int, float)):
        kind["数值"] += 1
    else:
        kind["文本"] += 1
        texts[str(v)] += 1
print("浓度列类型:", dict(kind))
print("非数值取值:", texts.most_common(12))

print("\n空浓度行:", sum(1 for r in rows if r[4] is None))
print("空文献名 :", sum(1 for r in rows if not r[0]))
print("空化合物 :", sum(1 for r in rows if not r[1]))
print("空 CAS   :", sum(1 for r in rows if not r[2]))
print("不同文献 :", len({r[0] for r in rows}))
print("不同样品 :", len({r[3] for r in rows}))

print("\n前 12 行:")
for r in rows[:12]:
    print("   %-34s %-26s %-12s %-16s %s"
          % (str(r[0])[:34], str(r[1])[:26], str(r[2])[:12], str(r[3])[:16], r[4]))

print("\n样例核对（源表里带 ± 和尾字母的）:")
for r in rows:
    if str(r[1]) in ("ethyl hexanoate", "ethyl butanoate") and r[0].startswith("2"):
        print("   %-24s %-14s %s" % (r[1][:24], str(r[3])[:14], r[4]))
wb.close()