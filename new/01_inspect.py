# -*- coding: utf-8 -*-
"""Step 1: inspect structure of the source Excel workbooks (read-only)."""
import sys, io, os
import pandas as pd
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

SRC = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
FILES = ['data总库.xlsx', 'data.xlsx']

out = []
for f in FILES:
    p = os.path.join(SRC, f)
    out.append('=' * 100)
    out.append(f'FILE: {f}   size={os.path.getsize(p):,} bytes')
    out.append('=' * 100)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    for s in wb.sheetnames:
        ws = wb[s]
        out.append(f'\n--- SHEET: {s!r}  dims={ws.max_row} x {ws.max_column}')
        # dump first 6 rows raw
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(6, ws.max_row), values_only=True)):
            cells = ['' if c is None else str(c)[:28] for c in row[:14]]
            out.append(f'   r{i+1}: ' + ' | '.join(cells))
    wb.close()

with open(os.path.join(os.path.dirname(__file__), 'output', '01_structure.txt'), 'w', encoding='utf-8') as fh:
    fh.write('\n'.join(out))
print('done', len(out), 'lines')
