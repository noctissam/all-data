# -*- coding: utf-8 -*-
"""Minimal probe: are rendered coordinates 1:1 with the slide?"""
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

p = Presentation()
p.slide_width = Inches(13.333)
p.slide_height = Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[6])
SPOTS = [(0.50, 0.50), (6.50, 0.50), (12.0, 0.50), (0.50, 6.50), (12.0, 6.50)]
for i, (l, t) in enumerate(SPOTS):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t),
                            Inches(0.80), Inches(0.40))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0, 0, 0) if i % 2 == 0 else RGBColor(200, 0, 0)
    sh.line.fill.background()
p.save("new/scratch/probe.pptx")
print("saved new/scratch/probe.pptx")
