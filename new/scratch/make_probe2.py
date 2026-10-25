# -*- coding: utf-8 -*-
"""Probe 2: does a 1200x675 pt (=16.667x9.375 in) slide render 1:1?"""
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

p = Presentation()
p.slide_width = Inches(16.667)
p.slide_height = Inches(9.375)
s = p.slides.add_slide(p.slide_layouts[6])
SPOTS = [(0.50, 0.50), (8.00, 0.50), (15.5, 0.50),
         (0.50, 8.50), (15.5, 8.50), (8.00, 8.50)]
for i, (l, t) in enumerate(SPOTS):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t),
                            Inches(0.80), Inches(0.40))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0, 0, 0)
    sh.line.fill.background()
p.save("new/scratch/probe2.pptx")
print("saved probe2")
