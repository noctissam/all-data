# -*- coding: utf-8 -*-
import os
from pptx import Presentation

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
p = Presentation(os.path.join(HERE, "数据清洗与关系挖掘_汇报.pptx"))
print(f"slides: {len(p.slides._sldIdLst)}   size: "
      f"{p.slide_width / 914400:.2f} x {p.slide_height / 914400:.2f} in")
for i, s in enumerate(p.slides, 1):
    titles = [sh.text_frame.text.split("\n")[0] for sh in s.shapes
              if sh.has_text_frame and sh.text_frame.text.strip()]
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    tabs = sum(1 for sh in s.shapes if sh.has_table)
    t = titles[0][:44] if titles else "(图片/图形页)"
    print(f"{i:>2}. {t:<46} [图 {pics}  表 {tabs}]")
