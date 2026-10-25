# -*- coding: utf-8 -*-
"""Step 2: full dump of the two 'all data' master sheets (read-only)."""
import sys, io, os
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
SRC = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
out = []

for fname, sheet in [('data总库.xlsx', 'all data'), ('data.xlsx', 'all data')]:
    p = os.path.join(SRC, fname)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb[sheet]
    out.append('#' * 110)
    out.append(f'{fname} :: sheet={sheet!r}  {ws.max_row} x {ws.max_column}')
    out.append('#' * 110)
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    for j, h in enumerate(hdr):
        out.append(f'  col[{j:3d}] {h!r}')
    out.append('')
    out.append('--- ROWS ---')
    for i, r in enumerate(rows[1:], start=2):
        cells = ['' if c is None else str(c) for c in r]
        # trim trailing empties
        while cells and cells[-1] == '':
            cells.pop()
        out.append(f'r{i:3d}| ' + ' || '.join(cells))
    wb.close()

with open(os.path.join(os.path.dirname(__file__), 'output', '02_alldata_raw.txt'), 'w', encoding='utf-8') as fh:
    fh.write('\n'.join(out))
print('done', len(out))
