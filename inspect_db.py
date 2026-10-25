# -*- coding: utf-8 -*-
"""检查 data总库.xlsx 与 清洗后数据.xlsx 的结构,输出到 UTF-8 文件"""
import openpyxl

BASE = r"C:\Users\24252\Desktop\PPT\课题\26年\26年3月"
OUT = BASE + r"\inspect_output.txt"

def fmt(v, w=30):
    if v is None:
        return ""
    return str(v)[:w]

with open(OUT, "w", encoding="utf-8") as f:
    for fname in ["data总库.xlsx", "清洗后数据.xlsx"]:
        path = BASE + "\\" + fname
        f.write("=" * 100 + "\n")
        f.write("FILE: " + fname + "\n")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        f.write("Sheet names: " + str(wb.sheetnames) + "\n")
        for ws in wb.worksheets:
            f.write(f"\n--- Sheet: {ws.title!r}  max_row={ws.max_row} max_col={ws.max_column} ---\n")
            # 打印前3行
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), 1):
                cells = [fmt(c, 25) for c in row]
                f.write(f"  R{i}: {cells}\n")
        wb.close()
        f.write("\n\n")

print("DONE -> inspect_output.txt")
